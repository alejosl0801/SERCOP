# -*- coding: utf-8 -*-
"""Lectura de anexos de procesos SERCOP: PDF (con OCR si está escaneado), Word (.docx/.doc/.rtf/.odt),
Excel e imágenes. Devuelve siempre lista de (pagina_o_hoja, texto). Nunca lanza."""
import shutil
import subprocess
import tempfile
from pathlib import Path

SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
MIN_CHARS_PAGINA = 60  # menos que esto en una página de PDF = probablemente escaneada
_ocr = None


def _motor_ocr():
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    return _ocr


def _ocr_imagen(png_bytes):
    res, _ = _motor_ocr()(png_bytes)
    if not res:
        return ""
    # res: [[caja, texto, confianza], ...] ordenado de arriba a abajo
    filas = sorted(res, key=lambda r: (round(r[0][0][1] / 12), r[0][0][0]))
    return "\n".join(r[1] for r in filas)


def _pdf(ruta):
    import fitz
    out = []
    d = fitz.open(str(ruta))
    for i, p in enumerate(d):
        t = p.get_text()
        if len(t.strip()) < MIN_CHARS_PAGINA:
            try:
                ocr = _ocr_imagen(p.get_pixmap(dpi=200).tobytes("png"))
                if len(ocr.strip()) > len(t.strip()):
                    t = "[OCR] " + ocr
            except Exception as e:
                t = t + f"\n[OCR FALLÓ: {e}]"
        out.append((i + 1, t))
    return out


def _docx(ruta):
    import docx
    d = docx.Document(str(ruta))
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            vistos = []
            for c in r.cells:
                tx = c.text.strip().replace("\n", " ")
                if tx and tx not in vistos:
                    vistos.append(tx)
            if vistos:
                partes.append(" | ".join(vistos))
    for sec in d.sections:
        for p in list(sec.header.paragraphs) + list(sec.footer.paragraphs):
            if p.text.strip():
                partes.append(p.text)
    return [(1, "\n".join(partes))]


def _convertir_con_libreoffice(ruta, formato):
    """.doc/.rtf/.odt/.xls -> docx/xlsx usando LibreOffice sin ventana."""
    if not SOFFICE.exists():
        raise RuntimeError("LibreOffice no está instalado")
    tmp = Path(tempfile.mkdtemp(prefix="sercop_conv_"))
    try:
        subprocess.run([str(SOFFICE), "--headless", "--convert-to", formato, "--outdir", str(tmp), str(ruta)],
                       capture_output=True, timeout=180, creationflags=0x08000000)
        salida = list(tmp.glob("*." + formato))
        if not salida:
            raise RuntimeError("LibreOffice no produjo salida")
        return salida[0], tmp
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise


def _xlsx(ruta):
    import openpyxl
    wb = openpyxl.load_workbook(str(ruta), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets[:8]:
        filas = []
        for row in ws.iter_rows(max_row=600, values_only=True):
            cel = [str(c) for c in row if c is not None and str(c).strip()]
            if cel:
                filas.append(" | ".join(cel))
        out.append((ws.title, "\n".join(filas)))
    return out


def _zip(ruta):
    import zipfile
    if ruta.suffix.lower() != ".zip":
        return [(0, "[FORMATO NO LEGIBLE AUTOMÁTICAMENTE: " + ruta.suffix + "]")]
    tmp = Path(tempfile.mkdtemp(prefix="sercop_zip_"))
    out = []
    try:
        with zipfile.ZipFile(ruta) as z:
            z.extractall(tmp)
        for f in sorted(tmp.rglob("*")):
            if f.is_file():
                for pag, txt in leer(f):
                    out.append((f"{f.name} · {pag}", txt))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return out or [(0, "[ZIP VACÍO]")]


def leer(ruta):
    ruta = Path(ruta)
    ext = ruta.suffix.lower()
    try:
        if ext in (".zip", ".rar", ".7z"):
            return _zip(ruta)
        if ext == ".pdf":
            return _pdf(ruta)
        if ext in (".docx", ".dotx"):
            return _docx(ruta)
        if ext in (".doc", ".rtf", ".odt", ".dot"):
            conv, tmp = _convertir_con_libreoffice(ruta, "docx")
            try:
                return _docx(conv)
            finally:
                shutil.rmtree(tmp, ignore_errors=True)
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(ruta)
        if ext in (".xls", ".ods"):
            conv, tmp = _convertir_con_libreoffice(ruta, "xlsx")
            try:
                return _xlsx(conv)
            finally:
                shutil.rmtree(tmp, ignore_errors=True)
        if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"):
            return [(1, "[OCR] " + _ocr_imagen(ruta.read_bytes()))]
        if ext in (".txt", ".csv"):
            return [(1, ruta.read_text(encoding="utf-8", errors="replace"))]
    except Exception as e:
        return [(0, f"[ERROR AL LEER: {e}]")]
    return [(0, "[FORMATO NO LEGIBLE AUTOMÁTICAMENTE]")]
