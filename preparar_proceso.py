# -*- coding: utf-8 -*-
"""Prepara un proceso SERCOP de punta a punta.

1. Entra al portal, descarga TODOS los anexos del proceso.
2. Lee cada línea de cada anexo (PDF, Word, Excel) y del propio portal.
3. Detecta requisitos y señales de alerta (notario, juramento, garantías, visitas, anexos obligatorios...).
4. Arma el borrador de proforma (ítems con CPC y descripción oficiales, precios según politica_sercop.json).
5. Lo firma con FirmaEC si todo tiene precio, y deja un resumen HTML con alertas y pendientes.

Uso:  py preparar_proceso.py NIC-xxxx [--estado "En Curso"|Finalizada] [--salida CARPETA] [--no-firmar]
      py preparar_proceso.py --cola      (procesa los códigos encolados por el vigilante)
"""
import argparse
import html
import json
import os
import re
import shutil
import sys
import traceback
import unicodedata
from datetime import datetime
from pathlib import Path

for _s in ("stdout", "stderr"):
    if getattr(sys, _s) is None:  # pythonw: sin consola
        setattr(sys, _s, open(os.devnull, "w", encoding="utf-8"))
    else:
        getattr(sys, _s).reconfigure(encoding="utf-8", errors="replace")

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

ESTADO_DIR = Path(r"C:\Users\HP\Dropbox\Previfuego\ESTADO")
RUC_SRI = Path(r"C:\Users\HP\Dropbox\PROFORMAS PAGINA WEB\NO APROBADOS\005511_ThePoint\requisitos\Certificado_RUC_Lopez.pdf")
STOCK_IMP = Path(r"C:\Users\HP\Dropbox\PyroShield\STOCK IMP.xlsx")
STOCK_REAL = Path(r"C:\Users\HP\Dropbox\PyroShield\STOCK REAL PYROSHIELD.xlsx")
POLITICA = BASE / "politica_sercop.json"
COLA = BASE / "cola_preparar.jsonl"
LOG = BASE / "preparar.log"

NCO_URL = "https://compraspublicas.gob.ec/ProcesoContratacion/compras/NCO/FrmNCOListado.cpe"
PORTAL = "https://compraspublicas.gob.ec/ProcesoContratacion/compras/"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36"
ESTADOS = {"En Curso": "384", "Finalizada": "476"}


def log(msg):
    try:
        with open(LOG, "a", encoding="utf-8") as f:
            f.write(f"{datetime.now():%Y-%m-%d %H:%M:%S}  {msg}\n")
    except Exception:
        pass


def norm(s):
    """minúsculas sin tildes, conservando la longitud para poder citar el texto original"""
    t = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").lower()
    return t if len(t) == len(s) else s.lower()


def slug(s, n=40):
    s = re.sub(r"[^A-Za-z0-9]+", "_", "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"))
    return s.strip("_")[:n] or "archivo"


# ─────────────────────────── 1. PORTAL ────────────────────────────────────────

def leer_portal(codigo, estado, carpeta_anexos):
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            page = browser.new_context(locale="es-EC", user_agent=UA).new_page()
            href = None
            for term in ("extintor", "recarga", "incendio"):
                for _ in range(3):
                    try:
                        page.goto(NCO_URL, wait_until="domcontentloaded", timeout=60000)
                        page.wait_for_selector("table tbody tr", timeout=45000)
                        page.wait_for_timeout(2500)
                        selects = page.locator("select").all()
                        if estado != "En Curso" and len(selects) > 2:
                            selects[2].select_option(ESTADOS[estado])
                            page.wait_for_timeout(1500)
                        box = page.locator("input[type='search']").first
                        box.fill(term)
                        box.press("Enter")
                        page.wait_for_timeout(3500)
                        for sel in page.locator("select").all():
                            vals = [o.get_attribute("value") for o in sel.locator("option").all()]
                            if "100" in vals:
                                sel.select_option("100")
                                page.wait_for_timeout(2500)
                                break
                        for row in page.locator("table tbody tr").all():
                            tds = row.locator("td").all()
                            if len(tds) > 1 and codigo in tds[1].inner_text():
                                href = row.locator("a").first.get_attribute("href")
                                break
                        if href:
                            break
                    except Exception as e:
                        log(f"buscar {term}: {str(e).splitlines()[0][:120]}")
                if href:
                    break
            if not href:
                raise RuntimeError(f"No encontré {codigo} en el listado ({estado}).")

            url = PORTAL + "NCO/" + href.split("/")[-1]
            for intento in range(4):
                try:
                    page.goto(url, wait_until="domcontentloaded", timeout=90000)
                    break
                except Exception as e:
                    log(f"detalle intento {intento + 1}: {str(e).splitlines()[0][:100]}")
                    if intento == 3:
                        raise
                    page.wait_for_timeout(5000)
            page.wait_for_timeout(3000)
            texto = page.inner_text("body")
            filas = page.evaluate("""() => [...document.querySelectorAll('tr')]
                .map(tr => [...tr.children].map(td => td.innerText.replace(/\\s+/g,' ').trim()))
                .filter(r => r.length >= 6 && /^\\d{6,10}$/.test(r[1]))""")
            links = page.evaluate("""() => [...document.querySelectorAll('a')]
                .filter(a => (a.getAttribute('href')||'').includes('ExeGENBajarArchivoGeneral'))
                .map(a => ({href: a.getAttribute('href'),
                            label: ((a.closest('tr') ? a.closest('tr').innerText : a.innerText) || '').replace(/\\s+/g,' ').trim()}))""")

            carpeta_anexos.mkdir(parents=True, exist_ok=True)
            anexos = []
            for i, l in enumerate(links, 1):
                full = PORTAL + "GE/" + l["href"].split("GE/")[-1]
                r = page.request.get(full, timeout=90000)
                cd = r.headers.get("content-disposition", "")
                m = re.search(r'filename="?([^";]+)', cd)
                ext = Path(m.group(1)).suffix.lower() if m else ".bin"
                etiqueta = l["label"][:70]
                ruta = carpeta_anexos / f"{i:02d}_{slug(etiqueta, 28)}{ext}"
                ruta.write_bytes(r.body())
                anexos.append({"n": i, "etiqueta": etiqueta, "ruta": ruta, "ext": ext, "bytes": len(r.body()),
                               "http": r.status})
            return {"url": url, "texto": texto, "filas": filas, "anexos": anexos}
        finally:
            browser.close()


def val_despues(texto, etiqueta):
    m = re.search(re.escape(etiqueta) + r"\s*:?\s*\n*\s*([^\n]+)", texto)
    return m.group(1).strip() if m else ""


def funcionario_de(texto):
    m = re.search(r"Funcionario Encargado\s+Nombre:?\s*([^\n]+)", texto) or re.search(r"Nombre\s*:\s*([^\n]+)", texto)
    return m.group(1).strip() if m else ""


def campos_portal(texto):
    forma = re.search(r"Forma de Pago\s*\n(.*?)\n\s*Documentos Anexos", texto, re.S)
    return {
        "entidad": val_despues(texto, "Nombre Entidad"),
        "tipo": val_despues(texto, "Tipo de necesidad"),
        "objeto": val_despues(texto, "Objeto de compra"),
        "estado": val_despues(texto, "Estado de la necesidad"),
        "publicacion": val_despues(texto, "Fecha de Publicación de la Necesidad"),
        "limite": val_despues(texto, "Fecha Límite para la entrega de Proformas"),
        "funcionario": funcionario_de(texto),
        "email": val_despues(texto, "Correo Electrónico"),
        "provincia": val_despues(texto, "Provincia"),
        "canton": val_despues(texto, "Cantón"),
        "parroquia": val_despues(texto, "Parroquia"),
        "direccion": val_despues(texto, "Dirección"),
        "forma_pago_portal": re.sub(r"\s+", " ", forma.group(1)).strip() if forma else "",
    }


# ─────────────────────────── 2. LECTURA DE ANEXOS ─────────────────────────────

def leer_anexo(ruta):
    """Delegado en lector_anexos (PDF+OCR, Word/.doc vía LibreOffice, Excel, imágenes)."""
    import lector_anexos
    return lector_anexos.leer(ruta)


# ─────────────────────────── 3. REQUISITOS Y ALERTAS ──────────────────────────

SECCIONES = {
    "plazo": r"plazo\s+de\s+ejecuci[oó]n|plazo\s+de\s+entrega",
    "pago": r"forma\s+y\s+condiciones\s+de\s+pago|forma\s+de\s+pago",
    "vigencia": r"vigencia\s+de\s+(?:la\s+)?(?:oferta|proforma)",
    "garantia": r"garant[ií]a\s+t[eé]cnica",
    "lugar": r"lugar\s+de\s+entrega|lugar\s+de\s+ejecuci[oó]n",
    "multas": r"\bmultas?\b",
    "obligaciones": r"obligaciones\s+del\s+contratista",
    "personal": r"personal\s+t[eé]cnico",
    "equipo": r"equipo\s+o\s+maquinaria",
    "administrador": r"administrador\s+de\s+la\s+orden",
}
CORTE = re.compile(r"\s\d{1,2}\s?\.\s?-\s*[A-ZÁÉÍÓÚÑ]{3,}")

# (clave, nivel, título, patrón sin tildes, qué hacer)
REGLAS = [
    ("notarial", "CRÍTICO", "Pide documento NOTARIZADO / ante notario",
     r"notari|escritura publica|protoliz", "Exige notaría: no se resuelve con firma electrónica sola. Decidir si participa."),
    ("juramento", "CRÍTICO", "Pide DECLARACIÓN JURAMENTADA",
     r"juramenta|bajo juramento", "Declaración bajo juramento a nombre de Alejandro; no incluirla sin su autorización."),
    ("garantias", "CRÍTICO", "Pide garantías o pólizas",
     r"garantia bancaria|poliza de (?:seguro|fiel|buen)|fiel cumplimiento|buen uso del anticipo|garantia de cumplimiento",
     "Garantía bancaria o póliza: costo y trámite previos."),
    ("visita", "CRÍTICO", "Exige visita técnica o inspección previa",
     r"visita tecnica (?:obligatoria|previa)|visita (?:tecnica )?(?:es )?obligatoria|obligatori[oa][^.]{0,40}visita|visita previa|inspeccion previa (?:del|de la|de las|al) (?:sitio|local|lugar|instalacion|inmueble|edificio)|inspeccion previa obligatoria|reconocimiento (?:del )?(?:sitio|lugar)",
     "Visita obligatoria antes de ofertar: confirmar fecha y si es requisito de validez."),
    ("anexos", "ATENCIÓN", "Anexos / formularios que deben presentarse",
     r"\banexo\s*(?:no\.?\s*)?(?:\d+|[ivx]+)\b|check ?list|formulario", "Revisar cada anexo y quién lo firma."),
    ("capacitacion", "ATENCIÓN", "Capacitaciones o charlas incluidas en el precio",
     r"capacitaci|charla", "Costo de horas técnicas no cotizado en el precio por libra."),
    ("incluidos", "ATENCIÓN", "Repuestos, soportes, pintura, etiquetas o provisionales incluidos",
     r"reemplazo de|cambio de (?:accesorios|repuestos)|repuestos?\b|soporte|repint|etiquet|sin costo adicional|provisional",
     "Cuesta dinero y suele ir DENTRO del precio: ajustar tarifa si hace falta."),
    ("logistica", "ATENCIÓN", "Logística especial (en sitio, horarios, traslados)",
     r"in situ|en sitio|retiro de|traslado|transporte|galpon|horario de|\d{1,2}h\d{2}\s*(?:a|-)\s*\d{1,2}h\d{2}",
     "Recarga en sitio exige equipo móvil y personal en horario fijo."),
    ("certificados", "ATENCIÓN", "Certificaciones, permisos o habilitantes",
     r"certificad[oa]s? (?:de|del|por|bomberos)|\brup\b|permiso de funcionamiento|patente|licencia|cuerpo de bomberos|iso ?\d{4,5}",
     "Confirmar que se tienen vigentes y adjuntarlos."),
    ("experiencia", "ATENCIÓN", "Experiencia o contratos similares",
     r"experiencia|contratos? similares?", "Verificar si hay mínimo de experiencia."),
    ("personal", "ATENCIÓN", "Personal, equipos, seguros o uniformes mínimos",
     r"personal tecnico|tecnico certificado|equipo minimo|maquinaria|uniforme|\bepp\b|seguro de accidentes|afiliacion al iess",
     "Requisitos de recursos mínimos."),
    ("pago", "INFO", "Condiciones de pago",
     r"anticipo|contra informe|informe de satisfaccion|contra entrega|a credito|\b\d{2} dias (?:plazo|despues)",
     "Define cuándo se cobra."),
    ("multas", "INFO", "Multas por atraso", r"multa|1 ?x ?1000|1 ?por ?mil|penalidad", "Calcular riesgo."),
    ("entregables", "INFO", "Entregables obligatorios (informes, actas, fotos, cronograma)",
     r"informe tecnico|acta de entrega|evidencia fotografica|cronograma", "Preparar formatos."),
]


def aplanar(s):
    return re.sub(r"\s+", " ", s).strip()


def extraer_desde(plano, m, max_chars=420):
    resto = plano[m.end(): m.end() + 1500]
    c = CORTE.search(resto)
    if c:
        resto = resto[: c.start()]
    resto = re.sub(r"\(¿[^)]*\)", " ", resto)
    resto = re.sub(r"^\s*[A-ZÁÉÍÓÚÑ0-9/.,()¿? -]{8,}(?=[A-ZÁÉÍÓÚÑ][a-záéíóúñ])", "", aplanar(resto))
    frases = [f.strip() for f in re.split(r"(?<=[.;])\s+", resto) if f.strip()]
    frases = [f for f in frases if not re.match(r"^INDICAR\b", f) and len(re.sub(r"\W", "", f)) > 3]
    out = ""
    for f in frases:
        if len(out) + len(f) > max_chars:
            break
        out += (" " if out else "") + f
    return out.strip(" .:-") + ("." if out and not out.endswith(".") else "")


VALIDADORES = {
    "plazo": lambda s: bool(re.search(r"\d+\s*d[ií]as", s, re.I)),
    "pago": lambda s: bool(re.search(r"pago|%|contra|anticipo|informe|factura|entrega|cr[eé]dito", s, re.I)) and not s.startswith("Provincia"),
    "garantia": lambda s: bool(re.search(r"a[nñ]o|mes|garant|meses", s, re.I)),
    "lugar": lambda s: len(s) > 25 and not s.startswith("Provincia"),
}


OBLIGACION = re.compile(
    r"\b(deber[aá]n?|debe|obligatori[oa]s?|obligad[oa]s?|exig|requiere[n]?|se requiere|presentar|adjuntar|entregar[aá]?n?|"
    r"incluir|incluye[n]?|contar[aá]n? con|cumplir|certificad|sin costo|a su cargo|por cuenta del|prohibid|no se aceptar)",
    re.I)
PLANTILLA = re.compile(r"\b(COLOCAR|COMPLETAR|INDICAR|ESTABLECER|SELECCIONAR)\b|\(Caso de|n[uú]mero de d[ií]as|xxxx", re.I)


RUIDO_LEGAL = re.compile(r"\b(Ley|Reglamento|Art[ií]culo|Resoluci[oó]n|Servicio Nacional de Contrataci[oó]n|audiencias?|herramienta inform[aá]tica)\b")


PLACEHOLDER = re.compile(r"de conformidad a lo requerido|conforme a lo requerido en los t[eé]rminos|\(de conformidad", re.I)


def obligaciones(corpus, maximo=150):
    """Red de seguridad independiente de las palabras clave: TODA frase de los anexos de la entidad que
    impone algo al proveedor (deberá, obligatorio, presentar, incluir…). El TDR va primero; se omiten citas legales genéricas."""
    def prioridad(fuente):
        f = norm(fuente)
        return 0 if ("terminos de referencia" in f or "tdr" in f or "especificacion" in f) else 1
    por_doc = {}
    for fuente, pag, texto in corpus:
        if fuente.startswith("PORTAL") or texto.startswith("["):
            continue
        por_doc.setdefault(fuente, []).append((pag, aplanar(texto)))
    vistas, out = set(), []
    for fuente in sorted(por_doc, key=prioridad):
        for pag, texto in por_doc[fuente]:
            for f in re.split(r"(?<=[.;:])\s+", texto):
                f = f.strip()
                if not (25 <= len(f) <= 600) or not OBLIGACION.search(f) or PLANTILLA.search(f) or RUIDO_LEGAL.search(f):
                    continue
                h = norm(f)[:80]
                if h in vistas:
                    continue
                vistas.add(h)
                out.append({"fuente": fuente.split(" [")[0], "pagina": pag, "texto": f})
    return out[:maximo]


def analizar(corpus, portal):
    """corpus: lista de (fuente, pagina, texto). Devuelve campos clave y alertas con evidencia.
    Los anexos de la entidad (TDR) tienen prioridad sobre el texto del portal."""
    por_fuente = {}
    for fuente, pag, texto in corpus:
        por_fuente.setdefault(fuente, []).append(aplanar(texto))
    adjuntos = [" ".join(v) for k, v in por_fuente.items() if not k.startswith("PORTAL")]
    planos = adjuntos + [" ".join(v) for k, v in por_fuente.items() if k.startswith("PORTAL")]

    def buscar(patron, validar=None):
        primero = ""
        for plano in planos:
            for m in re.finditer(patron, plano, re.I):
                s = extraer_desde(plano, m)
                if s and PLACEHOLDER.search(s):
                    continue
                if s:
                    primero = primero or s
                    if validar is None or validar(s):
                        return s
        return primero if validar is None else ""

    campos = {k: buscar(pat, VALIDADORES.get(k)) for k, pat in SECCIONES.items()}

    def regex_global(rx):
        for plano in planos:
            m = re.search(rx, plano, re.I)
            if m:
                return m
        return None

    vig = regex_global(r"vigencia\s+de\s+(?:la\s+)?(?:oferta|proforma)\D{0,60}?(\d{1,3})\s*d[ií]as")
    campos["vigencia_dias"] = int(vig.group(1)) if vig else None
    gar = regex_global(r"garant[ií]a\s+t[eé]cnica\D{0,80}?(un|una|1|dos|2|tres|3)\s+a[nñ]os?")
    campos["garantia_anios"] = {"un": 1, "una": 1, "1": 1, "dos": 2, "2": 2, "tres": 3, "3": 3}.get(gar.group(1).lower()) if gar else None
    pz = re.search(r"(\d{1,3})\s*d[ií]as\s*(calendario|h[aá]biles|laborables)?", campos.get("plazo", ""), re.I)
    campos["plazo_dias"] = (int(pz.group(1)), (pz.group(2) or "").lower()) if pz else None

    campos["obligaciones"] = obligaciones(corpus)

    import extractor_condiciones as ec
    cond = ec.resolver(ec.extraer([c for c in corpus if not c[0].startswith("PORTAL")]))
    campos["condiciones"] = cond
    if cond["vigencia"]:
        e = cond["vigencia"]["elegido"]
        campos["vigencia_dias"] = e["dias"]
    if cond["plazo"]:
        e = cond["plazo"]["elegido"]
        campos["plazo_dias"] = (e["dias"], e["tipo"])
        campos["plazo"] = e["cita"]
    if cond["garantia"]:
        campos["garantia_meses"] = cond["garantia"]["elegido"]["dias"] // 30
    else:
        campos["garantia_meses"] = None

    alertas = []
    for clave, nivel, titulo, patron, accion in REGLAS:
        rx = re.compile(patron, re.I)
        evid = []
        cubierto = {}
        for fuente, pag, texto in corpus:
            plano = aplanar(texto)
            n = norm(plano)
            for m in rx.finditer(n):
                rangos = cubierto.setdefault((fuente, pag), [])
                if any(a <= m.start() <= b for a, b in rangos):
                    continue
                a, b = max(0, m.start() - 110), min(len(plano), m.end() + 150)
                rangos.append((a, b))
                evid.append({"fuente": fuente, "pagina": pag, "texto": plano[a:b], "ctx": plano[max(0, a - 800):b]})
                if len(evid) >= 3:
                    break
            if len(evid) >= 3:
                break
        if evid and clave in ("notarial", "juramento") and all(
                re.search(r"previo a la suscripcion|previos a la suscripcion|procedera unicamente|estrictamente necesaria|expresamente prevista en los pliegos|para suscribir (?:el )?contrato|adjudicad|adjudicatari|ganador|una vez (?:adjudic|selecc)", norm(e.get("ctx", e["texto"]))) for e in evid):
            nivel = "ATENCIÓN"
            titulo += " (solo si resultas adjudicado)"
            accion = "El TDR lo exige después de ganar, previo a la orden de compra: no impide ofertar. Evalúa si lo harías."
        if evid:
            alertas.append({"clave": clave, "nivel": nivel, "titulo": titulo, "accion": accion, "evidencia": evid})

    lim = campos.get("plazo_dias")
    if lim and lim[0] <= 5 and "calend" in (lim[1] or "calend"):
        alertas.insert(0, {"clave": "plazo_corto", "nivel": "ATENCIÓN", "titulo": f"Plazo de ejecución muy corto: {lim[0]} días",
                           "accion": "Confirmar que se puede cumplir.", "evidencia": [{"fuente": "TDR", "pagina": "", "texto": campos["plazo"]}]})
    ilegibles = [f for f, p, t in corpus if t.startswith("[ERROR") or t.startswith("[FORMATO") or (len(t.strip()) < 40 and not f.startswith("PORTAL"))]
    for f in sorted(set(ilegibles)):
        alertas.insert(0, {"clave": "ilegible", "nivel": "ATENCIÓN", "titulo": f"No pude leer el texto de: {f}",
                           "accion": "Puede ser un PDF escaneado: leerlo a mano.", "evidencia": []})
    nombres = {"plazo": "plazo de ejecución", "vigencia": "vigencia de la oferta", "garantia": "garantía técnica"}
    for campo, v in cond.items():
        if v and v["contradiccion"]:
            vistos_c, ev = set(), []
            for cd in sorted(v["candidatos"], key=lambda x: x["dias"]):
                k = (cd["valor"], cd["unidad"], cd["fuente"])
                if k not in vistos_c:
                    vistos_c.add(k)
                    ev.append({"fuente": cd["fuente"], "pagina": cd["pagina"], "texto": cd["cita"]})
            regla = "el MENOR (más exigente)" if campo == "plazo" else "el MAYOR"
            alertas.append({"clave": "contradiccion_" + campo, "nivel": "ATENCIÓN",
                            "titulo": f"Los documentos se contradicen en {nombres[campo]}",
                            "accion": f"Apliqué {regla} en el borrador; confirma con la entidad si hay duda.", "evidencia": ev[:4]})
    orden = {"CRÍTICO": 0, "ATENCIÓN": 1, "INFO": 2}
    alertas.sort(key=lambda a: orden[a["nivel"]])
    return campos, alertas


# ─────────────────────────── 4. PRECIOS, STOCK Y BORRADOR ─────────────────────

def cargar_politica():
    return json.loads(POLITICA.read_text(encoding="utf-8"))


def cargar_lista_precios():
    out = {}
    try:
        import openpyxl
        ws = openpyxl.load_workbook(STOCK_IMP, data_only=True, read_only=True)["Lista de Precios"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            cap, ag, _c, nuevos, _d, clientes, stock = (list(row) + [None] * 7)[:7]
            m = re.search(r"([\d.,]+)", str(cap or ""))
            if not (m and ag):
                continue
            out[(str(ag).strip().upper(), float(m.group(1).replace(",", ".")))] = {
                "nuevos": nuevos, "clientes": clientes, "stock": stock}
    except Exception as e:
        log(f"lista de precios: {e}")
    return out


def cargar_stock_real():
    out = {}
    try:
        import openpyxl
        ws = openpyxl.load_workbook(STOCK_REAL, data_only=True, read_only=True)["Stock Real"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            cod, nombre, _i, stock, costo, precio1 = (list(row) + [None] * 6)[:6]
            m = re.match(r"EXTINTOR ([\d.]+) LIBRAS - (GAS CARBONICO|POLVO QUIMICO SECO)\s*$", str(nombre or ""))
            if m:
                ag = "CO2" if m.group(2).startswith("GAS") else "PQS"
                out[(ag, float(m.group(1)))] = {"codigo": cod, "stock": stock, "costo": costo, "precio1": precio1}
    except Exception as e:
        log(f"stock real: {e}")
    return out


def tarifa_por_volumen(n_recargas, pol):
    for t in pol["recarga_tarifa_por_libra"]["tramos"]:
        if n_recargas <= t["hasta_extintores"]:
            return t["tarifa"]
    return pol["recarga_tarifa_por_libra"]["tramos"][-1]["tarifa"]


def construir_items(filas, pol, codigo=None, objeto=""):
    """filas del portal -> items para generar_proforma + auditoría de precios."""
    import generar_proforma as gp

    lista, stock = cargar_lista_precios(), cargar_stock_real()
    items, auditoria, pendientes = [], [], []
    n_recargas = 0

    try:
        ajuste = json.loads((BASE / "ajustes_procesos.json").read_text(encoding="utf-8")).get(codigo or "", {})
    except Exception:
        ajuste = {}
    excluir = [norm(x) for x in ajuste.get("excluir_items", [])]
    for r in filas:
        if excluir and any(x in norm(r[3]) for x in excluir):
            log(f"{codigo}: ítem excluido por ajuste: {r[3][:80]}")
            continue
        cpc, desc_of, unidad, cant = r[1], r[3], r[4] if len(r) > 4 else "Unidad", r[5] if len(r) > 5 else "1"
        cant_f = float(str(cant).replace(",", "")) if re.match(r"^[\d.,]+$", str(cant)) else 1.0
        det = re.sub(r"\bP\.?Q\.?S\.?\b", "PQS", desc_of.replace(",", "."), flags=re.I)
        tipo, agente, lbs = gp.detect_tipo(det), gp.detect_agente(det), gp.extract_lbs(det)
        interpretado = False
        if tipo == "desconocido" and "recarga" in norm(objeto) and "adquisicion" not in norm(objeto) and agente in ("PQS", "CO2") and lbs:
            tipo, interpretado = "recarga", True  # el objeto del proceso es un servicio de recarga
            det = "Recarga de " + det
        it = {"descripcion": det, "descripcion_oficial": desc_of, "cpc": cpc, "unidad": unidad or "Unidad",
              "cantidad": cant_f}
        info = {"descripcion": desc_of, "cantidad": cant_f, "tipo": tipo, "agente": agente, "lbs": lbs,
                "fuente": "", "precio": None, "alertas": []}
        nd = norm(desc_of)

        servicio = None
        for s in pol["servicios_con_precio_conocido"]:
            if s["clave"].startswith("recarga extintor agua") and tipo == "recarga" and "agua" in nd and re.search(r"2[.,]5", desc_of):
                servicio = s
            if s["clave"].startswith("mantenimiento de manometros") and "manometro" in nd and "mantenimiento" in nd:
                servicio = s
        if servicio:
            it["precio_unitario"] = servicio["precio"]
            info.update(precio=servicio["precio"], fuente=f"precedente: {servicio['fuente']}")
        elif tipo == "recarga" and agente in ("PQS", "CO2") and lbs:
            n_recargas += int(cant_f)
            info["fuente"] = "tarifa por libra (se define al final)"
        elif tipo == "adquisicion" and agente in ("PQS", "CO2") and lbs:
            ref_l, ref_s = lista.get((agente, lbs), {}), stock.get((agente, lbs), {})
            precio, fuente = None, ""
            for fte in pol["extintores_nuevos"]["orden_de_fuentes"]:
                if fte == "precios_json_sercop":
                    for e in gp.ADQ_CATALOG:
                        if e["agente"] == agente and e["capacidad_lbs"] == lbs:
                            precio, fuente = e["precio_unitario"], "lista SERCOP (precedente ATM Daule)"
                elif fte == "azur_clientes_nuevos" and ref_l.get("nuevos"):
                    precio, fuente = round(float(ref_l["nuevos"]), 2), "Azur: precio clientes nuevos"
                if precio:
                    break
            if precio:
                it["precio_unitario"] = precio
            info.update(precio=precio, fuente=fuente or "SIN PRECIO",
                        azur_codigo=ref_s.get("codigo"), azur_clientes=ref_l.get("clientes"),
                        azur_nuevos=ref_l.get("nuevos"), costo=ref_s.get("costo"), stock=ref_s.get("stock"))
            if ref_s.get("stock") is not None and ref_s["stock"] < cant_f:
                info["alertas"].append(f"Stock insuficiente: hay {ref_s['stock']} y piden {int(cant_f)}")
            if precio and ref_s.get("costo"):
                marg = (precio - float(ref_s["costo"])) / precio * 100
                info["margen_pct"] = round(marg, 1)
                if marg < pol["alertas"]["margen_minimo_nuevos_pct"]:
                    info["alertas"].append(f"Margen bajo: {marg:.1f}%")
            if not precio:
                pendientes.append(f"Sin precio: {desc_of}")
        else:
            info["fuente"] = "NO RECONOCIDO"
            pendientes.append(f"Ítem no reconocido, necesito tu precio: {desc_of}")
        if interpretado and tipo == "recarga":
            info["alertas"].append("Interpretado como RECARGA porque el objeto del proceso es 'mantenimiento y recarga'")
        items.append(it)
        auditoria.append(info)

    tarifa = tarifa_por_volumen(n_recargas, pol) if n_recargas else None
    return items, auditoria, pendientes, tarifa, n_recargas


def siguiente_numero():
    mx = 3774
    for d in ESTADO_DIR.iterdir() if ESTADO_DIR.exists() else []:
        m = re.match(r"^(\d{4})\s", d.name)
        if m and d.is_dir():
            mx = max(mx, int(m.group(1)))
    return mx + 1


def carpeta_de_proceso(codigo, entidad):
    for d in ESTADO_DIR.iterdir() if ESTADO_DIR.exists() else []:
        if d.is_dir() and (d / "PROCESO.txt").exists() and codigo in (d / "PROCESO.txt").read_text(encoding="utf-8", errors="ignore"):
            return d, int(re.match(r"^(\d+)", d.name).group(1))
    n = siguiente_numero()
    corto = re.sub(r"\b(GOBIERNO AUTONOMO DESCENTRALIZADO|EMPRESA PUBLICA MUNICIPAL DE|EMPRESA PUBLICA|DE LA|DEL|DE|EP)\b", " ",
                   re.sub(r"[^\w ]", " ", entidad.upper()))
    corto = " ".join(corto.split()[:4]) or "ENTIDAD"
    d = ESTADO_DIR / f"{n} {corto}"
    d.mkdir(parents=True, exist_ok=True)
    return d, n


# ─────────────────────────── 5. RESUMEN ───────────────────────────────────────

def faltan(lim_txt):
    try:
        lim = datetime.strptime(lim_txt[:19], "%Y-%m-%d %H:%M:%S")
    except Exception:
        return "?", None
    s = (lim - datetime.now()).total_seconds()
    if s <= 0:
        return "VENCIDO", lim
    d, h, m = int(s // 86400), int(s % 86400 // 3600), int(s % 3600 // 60)
    return (f"{d} d {h} h" if d else f"{h} h {m} min"), lim


def esc(x):
    return html.escape(str(x if x is not None else ""))


def escribir_resumen(carpeta, num, portal_campos, nco, campos, alertas, auditoria, pendientes, res, anexos, firmado, url):
    quedan, _ = faltan(portal_campos["limite"])
    crit = [a for a in alertas if a["nivel"] == "CRÍTICO"]
    aten = [a for a in alertas if a["nivel"] == "ATENCIÓN"]
    nivel = "ROJO" if crit else ("NARANJA" if aten or pendientes else "VERDE")
    color = {"ROJO": "#b00020", "NARANJA": "#d35400", "VERDE": "#1e7e34"}[nivel]

    def bloque_alerta(a):
        ev = "".join(f"<li><b>{esc(e['fuente'])}{(' p.' + str(e['pagina'])) if e['pagina'] else ''}:</b> …{esc(e['texto'])}…</li>" for e in a["evidencia"])
        return (f"<div class='al {a['nivel'][:3]}'><b>[{esc(a['nivel'])}] {esc(a['titulo'])}</b>"
                f"<div class='acc'>{esc(a['accion'])}</div><ul>{ev}</ul></div>")

    filas = ""
    for i, a in enumerate(auditoria, 1):
        extra = "; ".join(a["alertas"])
        info = ""
        if a.get("azur_codigo"):
            info = (f"Azur {esc(a['azur_codigo'])} · clientes ${a.get('azur_clientes')} · nuevos ${a.get('azur_nuevos')} · "
                    f"costo ${a.get('costo')} · stock {a.get('stock')} · margen {a.get('margen_pct', '—')}%")
        filas += (f"<tr><td>{i}</td><td>{esc(a['descripcion'])}</td><td>{a['cantidad']:g}</td>"
                  f"<td>{esc(a.get('fuente'))}</td><td class='n'>{esc(info)}"
                  f"{'<br><b style=color:#b00020>' + esc(extra) + '</b>' if extra else ''}</td></tr>")

    valores_portal = "\n".join(("%.2f" % a["precio"]) if a["precio"] is not None else "" for a in auditoria)
    tabla_portal = "\n".join(f"{i:>2}. {a['cantidad']:>6g} x {esc(a['descripcion'])}" for i, a in enumerate(auditoria, 1))
    h = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{esc(portal_campos['entidad'])} - resumen</title>
<style>body{{font-family:Segoe UI,Arial;margin:24px;max-width:1100px;color:#222}}
h1{{margin:0}} .banner{{background:{color};color:#fff;padding:14px 18px;border-radius:6px;margin:10px 0}}
.al{{border-left:6px solid #888;background:#f6f6f6;padding:8px 12px;margin:8px 0}}
.al.CRÍ{{border-color:#b00020;background:#fdecee}} .al.ATE{{border-color:#d35400;background:#fff3e8}}
.acc{{color:#555;font-size:13px;margin:2px 0}} ul{{margin:4px 0 0 18px;font-size:12.5px}}
table{{border-collapse:collapse;width:100%;font-size:13px}} td,th{{border:1px solid #bbb;padding:5px 7px;vertical-align:top}}
th{{background:#1a3a5c;color:#fff}} .n{{font-size:11.5px;color:#444}} pre{{background:#f4f4f4;padding:10px;font-size:12.5px}}
.kv td:first-child{{width:210px;font-weight:600;background:#f0f0f0}}</style></head><body>
<h1>{esc(portal_campos['entidad'])}</h1>
<div class="banner"><b>{nivel}</b> · {esc(codigo_de(nco))} · límite {esc(portal_campos['limite'])} · <b>faltan {esc(quedan)}</b></div>
<p>{esc(portal_campos['objeto'])}</p>

<button onclick="navigator.clipboard.writeText(document.getElementById('vals').value).then(()=>{{this.textContent='COPIADO ✔'}})"
 style="font-size:16px;padding:10px 18px;margin:6px 0;cursor:pointer">📋 Copiar valores unitarios (uno por línea, en orden del portal)</button>
<textarea id="vals" style="position:absolute;left:-9999px">{valores_portal}</textarea>
<details open><summary><b>Checklist antes de enviar</b></summary><ul style="list-style:none;padding-left:4px">
<li>☐ Leí las alertas críticas y de atención</li>
<li>☐ Revisé cantidades y valores unitarios contra el portal</li>
<li>☐ Plazo, vigencia ({esc(campos.get('vigencia_dias') or '?')} días) y garantía coinciden con el TDR</li>
<li>☐ Adjunto RUC y anexos que pida la entidad (ver 'Obligaciones')</li>
<li>☐ PDF firmado: <code>SUBIR_ESTE.pdf</code> (segunda etapa del portal)</li>
<li>☐ Acepté términos y pulsé el botón final de la segunda etapa</li>
<li>☐ Marqué la decisión: <code>py decidir.py {esc(codigo_de(nco))} enviado</code></li></ul></details>

<h2>Alertas ({len(crit)} críticas, {len(aten)} de atención)</h2>
{''.join(bloque_alerta(a) for a in alertas) or '<p>Sin alertas detectadas.</p>'}

<h2>Obligaciones que impone la entidad ({len(campos.get('obligaciones', []))} frases; lectura rápida, independiente de las alertas)</h2>
<details><summary>Ver todas</summary><ul>{''.join('<li><b>' + esc(o['fuente']) + (' p.' + str(o['pagina'])) + ':</b> ' + esc(o['texto']) + '</li>' for o in campos.get('obligaciones', []))}</ul></details>

<h2>Pendientes de tu decisión</h2>
<ul>{''.join('<li>' + esc(p) + '</li>' for p in pendientes) or '<li>Ninguno automático.</li>'}
<li>Tarifa de recarga aplicada: <b>${nco.get('tarifa_lb')}/lb</b> según el volumen ({res.get('n_recargas')} recargas). Política en politica_sercop.json.</li>
<li>Precio de extintores nuevos: ver columna de fuente en la tabla de ítems.</li>
<li>La carta LAFT (declaración bajo juramento) NO se incluyó.</li></ul>

<h2>Condiciones según el TDR</h2>
<table class="kv">
<tr><td>Vigencia de la oferta</td><td>{esc(campos.get('vigencia_dias') or 'NO ENCONTRADA (puse 30 por defecto)')} días</td></tr>
<tr><td>Plazo de ejecución</td><td>{esc(campos.get('plazo') or 'NO ENCONTRADO')}</td></tr>
<tr><td>Forma de pago</td><td>{esc(campos.get('pago') or portal_campos['forma_pago_portal'] or 'NO ENCONTRADA')}<br><i>Portal: {esc(portal_campos['forma_pago_portal'])}</i></td></tr>
<tr><td>Garantía técnica</td><td>{esc(campos.get('garantia') or 'NO ENCONTRADA')}</td></tr>
<tr><td>Lugar de entrega / ejecución</td><td>{esc(campos.get('lugar') or '')}<br><i>Portal: {esc(portal_campos['direccion'])}, {esc(portal_campos['canton'])}</i></td></tr>
<tr><td>Obligaciones del contratista</td><td>{esc(campos.get('obligaciones') or '—')}</td></tr>
<tr><td>Multas</td><td>{esc(campos.get('multas') or '—')}</td></tr>
<tr><td>Personal / equipo mínimo</td><td>{esc(campos.get('personal') or '—')} / {esc(campos.get('equipo') or '—')}</td></tr>
<tr><td>Administrador de la orden</td><td>{esc(campos.get('administrador') or '—')}</td></tr>
<tr><td>Contacto</td><td>{esc(portal_campos['funcionario'])} · {esc(portal_campos['email'])}</td></tr>
</table>

<h2>Ítems y precios (cada ítem conserva CPC y descripción de la entidad)</h2>
<table><tr><th>#</th><th>Descripción oficial</th><th>Cant.</th><th>Fuente del precio</th><th>Auditoría</th></tr>{filas}</table>
<p><b>Subtotal sin IVA: ${res.get('subtotal')}</b> · IVA ${res.get('iva')} · Total ${res.get('total')} ·
{'<b>proforma FIRMADA (borrador)</b>' if firmado else '<b style="color:#b00020">NO se firmó: hay ítems sin precio</b>'}</p>

<h2>Cómo enviarla en el portal (ya probado)</h2>
<ol><li>SERCOP → Necesidades de Contratación → buscar "extintor" → entrar al proceso → <b>Enviar Proforma</b> → login.</li>
<li>Escribir el valor unitario de cada ítem (columna "Valor Unitario Ofertado") y la vigencia: <b>{esc(nco.get('vigencia_dias'))}</b>. El total debe dar el subtotal de arriba.</li>
<li>Guardar. <b>Aparece una segunda etapa</b>: subir el PDF firmado (Archivo Obligatorio) y el RUC (Archivo Opcional), con descripción y botón "Subir".</li>
<li>Aceptar los términos y dar el botón final. Confirmar que el portal diga que quedó registrada.</li></ol>
<pre>{tabla_portal}</pre>

<h2>Archivos del proceso</h2>
<ul>{''.join('<li>' + esc(a['etiqueta']) + ' → ' + esc(a['ruta'].name) + ' (' + str(a['bytes']) + ' bytes)</li>' for a in anexos)}</ul>
<p><a href="{esc(url)}">Detalle en el portal</a> · Generado {datetime.now():%Y-%m-%d %H:%M}</p>
</body></html>"""
    out = carpeta / "00 RESUMEN Y ALERTAS.html"
    out.write_text(h, encoding="utf-8")
    return out, nivel, crit, aten


def codigo_de(nco):
    return nco.get("codigo", "")


# ─────────────────────────── 6. ORQUESTA ──────────────────────────────────────

def preparar(codigo, estado="En Curso", salida=None, firmar=True):
    import generar_proforma as gp

    pol = cargar_politica()
    tmp = Path(salida) if salida else None
    previo = (tmp or ESTADO_DIR) / "_descarga_tmp"
    portal = leer_portal(codigo, estado, previo / "anexos")
    pc = campos_portal(portal["texto"])
    log(f"{codigo}: portal leído, {len(portal['filas'])} ítems, {len(portal['anexos'])} anexos")

    if tmp:
        carpeta, num = tmp, 0
        carpeta.mkdir(parents=True, exist_ok=True)
    else:
        carpeta, num = carpeta_de_proceso(codigo, pc["entidad"])
    anexos_dir = carpeta / "ANEXOS DE LA ENTIDAD"
    anexos_dir.mkdir(exist_ok=True)
    for a in portal["anexos"]:
        dest = anexos_dir / a["ruta"].name
        shutil.copy2(a["ruta"], dest)
        a["ruta"] = dest
    shutil.rmtree(previo, ignore_errors=True)

    corpus = [("PORTAL (detalle del proceso)", 1, portal["texto"])]
    texto_completo = [f"##### PORTAL\n{portal['texto']}"]
    for a in portal["anexos"]:
        for pag, txt in leer_anexo(a["ruta"]):
            corpus.append((f"{a['etiqueta'] or a['ruta'].name} [{a['ruta'].name}]", pag, txt))
            texto_completo.append(f"\n##### {a['ruta'].name} · pág/hoja {pag}\n{txt}")
    (carpeta / "TEXTO COMPLETO DE LOS ANEXOS.txt").write_text("\n".join(texto_completo), encoding="utf-8")
    campos, alertas = analizar(corpus, pc)

    items, auditoria, pendientes, tarifa, n_rec = construir_items(portal["filas"], pol, codigo, pc["objeto"])
    vig = campos["vigencia_dias"] or pol["proforma"]["vigencia_por_defecto_dias"]
    if not campos["vigencia_dias"]:
        pendientes.append(f"No encontré la vigencia de la oferta en el TDR: puse {vig} días por defecto. Verifícala.")
    pd_ = campos.get("plazo_dias")
    plazo = (f"{pd_[0]} días{(' ' + pd_[1]) if pd_[1] else ''}, contados conforme a lo establecido en los términos de referencia."
             if pd_ else "Conforme a los términos de referencia de la entidad contratante.")
    if not pd_:
        pendientes.append("No encontré el plazo de ejecución en el TDR: revisar manualmente.")
    pago = campos["pago"] or pc["forma_pago_portal"] or "Conforme a los términos de referencia de la entidad contratante."
    gm = campos.get("garantia_meses")
    garantia = (f"Garantía técnica de {gm} meses por defectos de taller y funcionamiento, conforme al TDR."
                if gm and gm % 12 else
                f"Garantía técnica de {gm // 12} año{'s' if gm // 12 > 1 else ''} por defectos de taller y funcionamiento, conforme al TDR."
                if gm else "Garantía técnica conforme al TDR de la entidad contratante.")
    if not gm:
        pendientes.append("No encontré el periodo de garantía técnica en el TDR: revisa el texto de garantía de la proforma.")
    if re.search(r"describir", norm(campos.get("garantia", ""))):
        pendientes.append("El TDR pide describir cómo se responde por la garantía: revisa el texto de garantía de la proforma.")

    nco = {
        "id": "", "codigo": codigo, "tipo": pc["tipo"], "fecha_publicacion": pc["publicacion"],
        "provincia_canton": f"{pc['provincia']} - {pc['canton']}", "descripcion": pc["objeto"],
        "estado": pc["estado"], "fecha_limite": pc["limite"], "entidad": pc["entidad"],
        "funcionario": pc["funcionario"], "email": pc["email"], "direccion": pc["direccion"], "url_detalle": portal["url"],
        "plazo_entrega": plazo, "forma_pago": pago, "vigencia_dias": vig, "garantia": garantia,
        "lugar_entrega": campos["lugar"] or f"{pc['direccion']}, {pc['canton']}",
        "especificaciones": ["Se cumplirá íntegramente lo establecido en las especificaciones técnicas y términos de referencia "
                             "(TDR) de la entidad contratante para cada ítem, incluidos los entregables y obligaciones del contratista."],
        "incluir_carta": pol["proforma"]["incluir_carta_laft"], "items": items,
    }
    if tarifa is not None:
        nco["tarifa_lb"] = tarifa

    (carpeta / "PROCESO.txt").write_text(f"{codigo}\n{pc['entidad']}\n{portal['url']}\n", encoding="utf-8")
    (carpeta / "datos_proceso.json").write_text(json.dumps(nco, ensure_ascii=False, indent=2), encoding="utf-8")

    sin_precio = any(a.get("fuente") in ("NO RECONOCIDO", "SIN PRECIO") for a in auditoria)
    puede_firmar = firmar and not sin_precio and bool(gp.P12_PASS)
    if firmar and not gp.P12_PASS:
        pendientes.append("No pude descifrar la clave del certificado: la proforma quedó sin firmar.")
    res = {}
    try:
        res = gp.procesar_nco(nco, firmar=puede_firmar, forzar=True) or {}
    except Exception:
        log("procesar_nco falló:\n" + traceback.format_exc())
        pendientes.append("Error al generar el PDF de la proforma (ver preparar.log).")
    res["n_recargas"] = n_rec

    entidad_corta = re.sub(r"^\d+\s", "", carpeta.name)
    base_nombre = f"Presupuesto_ {entidad_corta}_Extintores_{num}" if num else f"Presupuesto_ {codigo}"
    firmado = False
    pdf = Path(res["pdf"]) if res.get("pdf") else None
    if pdf and pdf.exists():
        firmado = pdf.name.endswith("_signed.pdf")
        shutil.copy2(pdf, carpeta / "SUBIR_ESTE.pdf")
        shutil.copy2(pdf, carpeta / (("BORRADOR FIRMADO - " if firmado else "BORRADOR SIN FIRMAR - ") + base_nombre + ("-signed" if firmado else "") + ".pdf"))
    if RUC_SRI.exists():
        shutil.copy2(RUC_SRI, carpeta / "RUC SRI - LOPEZ MEJIA ALEJANDRO.pdf")
    for sobra in (BASE / f"proforma_{codigo}.html", BASE / f"proforma_{codigo}_unsigned.pdf"):
        sobra.unlink(missing_ok=True)

    (carpeta / "TABLA PARA EL PORTAL.txt").write_text(
        f"{pc['entidad']}\n{codigo}\nLímite: {pc['limite']}\nVigencia de la proforma: {vig}\n\n"
        + "\n".join(f"{i:>2}. cant {a['cantidad']:>6g} | unitario {a['precio'] if a['precio'] is not None else '(tarifa)'} | {a['descripcion']}"
                    for i, a in enumerate(auditoria, 1)), encoding="utf-8")

    # precio unitario efectivo de las recargas (lo calcula el generador) para la tabla del portal
    try:
        filas_gp = gp.build_filas(items, tarifa)
        for a, f in zip(auditoria, filas_gp):
            if a["precio"] is None and f.matched:
                a["precio"] = f.precio_unitario
        (carpeta / "TABLA PARA EL PORTAL.txt").write_text(
            f"{pc['entidad']}\n{codigo}\nLímite: {pc['limite']}\nVigencia de la proforma: {vig}\n\n"
            + "\n".join(f"{i:>2}. cant {a['cantidad']:>6g} | VALOR UNITARIO {a['precio'] if a['precio'] is not None else '???':>9} | {a['descripcion']}"
                        for i, a in enumerate(auditoria, 1))
            + f"\n\nSubtotal sin IVA: {res.get('subtotal')}\n", encoding="utf-8")
    except Exception:
        log("tabla portal: " + traceback.format_exc())

    html_path, nivel, crit, aten = escribir_resumen(carpeta, num, pc, nco, campos, alertas, auditoria, pendientes, res, portal["anexos"], firmado, portal["url"])
    quedan, _ = faltan(pc["limite"])
    resumen = {"codigo": codigo, "entidad": pc["entidad"], "limite": pc["limite"], "faltan": quedan, "nivel": nivel,
               "criticas": [a["titulo"] for a in crit], "atencion": [a["titulo"] for a in aten],
               "pendientes": pendientes, "subtotal": res.get("subtotal"), "total": res.get("total"),
               "firmado": firmado, "carpeta": str(carpeta), "html": str(html_path)}
    (carpeta / "resumen.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")
    log(f"{codigo}: listo ({nivel}) en {carpeta}")
    return resumen


def aviso_final(r, popup=True):
    from vigilante_sercop import mostrar_popup

    cuerpo = [f"ENTIDAD: {r['entidad']}", f"CODIGO: {r['codigo']}",
              f"LIMITE: {r['limite']}   (faltan {r['faltan']})",
              f"SUBTOTAL SIN IVA: ${r['subtotal']}   TOTAL CON IVA: ${r['total']}",
              "PROFORMA: " + ("FIRMADA, lista para revisar" if r["firmado"] else "SIN FIRMAR (faltan decisiones tuyas)"), ""]
    if r["criticas"]:
        cuerpo.append("ALERTAS CRITICAS (leer antes de seguir):")
        cuerpo += [f"  - {x}" for x in r["criticas"]]
        cuerpo.append("")
    if r["atencion"]:
        cuerpo.append("A TENER EN CUENTA:")
        cuerpo += [f"  - {x}" for x in r["atencion"][:7]]
        cuerpo.append("")
    if r["pendientes"]:
        cuerpo.append("PENDIENTE DE TI:")
        cuerpo += [f"  - {x}" for x in r["pendientes"][:6]]
        cuerpo.append("")
    cuerpo.append("CARPETA: " + r["carpeta"])
    cuerpo.append("Abre '00 RESUMEN Y ALERTAS.html' para ver cada requisito con su cita.")
    banner = {"ROJO": "BORRADOR LISTO - OJO: EXIGEN COSAS CRITICAS",
              "NARANJA": "BORRADOR LISTO - REVISA LAS ALERTAS",
              "VERDE": "BORRADOR LISTO PARA REVISAR"}[r["nivel"]]
    color = {"ROJO": "Red", "NARANJA": "DarkOrange", "VERDE": "DarkGreen"}[r["nivel"]]
    auto = [x for x in r["criticas"] if "NOTARIZADO" in x or "JURAMENTADA" in x]
    if auto:
        import decisiones
        decisiones.registrar(r["codigo"], "descarte_auto", "; ".join(auto))
        banner = "DESCARTADO AUTOMATICO: " + ("pide notarizacion" if any("NOTARIZADO" in x for x in auto) else "pide declaracion juramentada")
        color = "Goldenrod"
        cuerpo.insert(0, "Lo descarté porque exige algo que decidiste no hacer. Si quieres participar igual:\n"
                         f"  py decidir.py {r['codigo']} participa\n")
    if popup:
        mostrar_popup(banner, color, "\n".join(cuerpo), url2=r.get("html", ""))
    try:
        from vigilante_sercop import push_celular
        push_celular(banner, "\n".join(cuerpo)[:900])
    except Exception:
        log("push celular: " + traceback.format_exc())
    return banner, color, bool(auto)


def _ventanas_borrador():
    try:
        return bool(cargar_politica().get("avisos", {}).get("ventana_al_terminar_borrador", False))
    except Exception:
        return False


def _aviso_silencioso(resultados, fallos):
    """Sin ventana: deja el resumen en ALERTAS.txt y lo manda al celular."""
    from vigilante_sercop import push_celular
    L = []
    for r, _, auto in resultados:
        est = "DESCARTADO AUTO" if auto else {"ROJO": "CRITICO", "NARANJA": "REVISAR", "VERDE": "OK"}[r["nivel"]]
        L.append(f"[{est}] {r['entidad']} · faltan {r['faltan']} · ${r['total']} · {'FIRMADA' if r['firmado'] else 'SIN FIRMAR'}")
    for l, e in fallos:
        L.append(f"[NO PUDE PREPARAR] {l[:100]}")
    texto = "\n".join(L)
    with open(BASE / "ALERTAS.txt", "a", encoding="utf-8") as fh:
        fh.write(f"===== BORRADORES {datetime.now():%Y-%m-%d %H:%M} =====\n{texto}\n\n")
    push_celular("Borradores listos" if not fallos else "Borradores (con fallos)", texto[:900])
    log("aviso silencioso: " + texto.replace("\n", " | ")[:300])


def resumen_unico(resultados, fallos):
    """UNA sola ventana para todo lo que terminó en esta tanda (en vez de una por proceso)."""
    from vigilante_sercop import mostrar_popup
    orden = {"ROJO": 0, "NARANJA": 1, "VERDE": 2}
    resultados = sorted(resultados, key=lambda x: (x[2], orden[x[0]["nivel"]]))
    lineas = []
    for r, _, auto in resultados:
        estado = "DESCARTADO AUTOMATICO" if auto else {"ROJO": "CRITICO", "NARANJA": "REVISAR", "VERDE": "OK"}[r["nivel"]]
        lineas.append(f"[{estado}] {r['entidad']}\n   {r['codigo']}  ·  faltan {r['faltan']}  ·  "
                      f"total ${r['total']}  ·  {'FIRMADA' if r['firmado'] else 'SIN FIRMAR'}")
        for x in r["criticas"][:3]:
            lineas.append(f"   ! {x}")
        lineas.append(f"   {r['carpeta']}")
        lineas.append("")
    for l, e in fallos:
        lineas.append(f"[NO PUDE PREPARAR] {l[:120]}\n   {e.strip().splitlines()[-1][:160]}\n")
    activos = [x for x in resultados if not x[2]]
    peor = min((orden[x[0]["nivel"]] for x in activos), default=3)
    color = "Red" if peor == 0 or fallos else "DarkOrange" if peor == 1 else "DarkGreen" if peor == 2 else "Goldenrod"
    n = len(resultados) + len(fallos)
    banner = "BORRADOR LISTO" if n == 1 else f"{n} BORRADORES LISTOS"
    lineas.append("Abre 00 RESUMEN Y ALERTAS.html de cada carpeta antes de decidir.")
    html = next((r["html"] for r, _, a in resultados if not a), resultados[0][0]["html"] if resultados else "")
    mostrar_popup(banner, color, "\n".join(lineas), url2=html)


def procesar_cola():
    if not COLA.exists():
        return
    lineas = [l.lstrip("﻿") for l in COLA.read_text(encoding="utf-8-sig").splitlines() if l.strip()]
    COLA.write_text("", encoding="utf-8")
    resultados, fallos = [], []
    for l in lineas:
        for intento in (1, 2):
            try:
                job = json.loads(l)
                r = preparar(job["codigo"], job.get("estado", "En Curso"), job.get("salida"), job.get("firmar", True))
                banner, color, auto = aviso_final(r, popup=False)
                resultados.append((r, banner, auto))
                break
            except Exception:
                err = traceback.format_exc()
                log(f"FALLO preparar (intento {intento}):\n" + err)
                if intento == 2:
                    fallos.append((l, err))
    if resultados or fallos:
        try:
            if _ventanas_borrador():
                resumen_unico(resultados, fallos)
            else:
                _aviso_silencioso(resultados, fallos)
        except Exception:
            log("resumen único: " + traceback.format_exc())


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("codigo", nargs="?")
    ap.add_argument("--estado", default="En Curso", choices=list(ESTADOS))
    ap.add_argument("--salida")
    ap.add_argument("--no-firmar", action="store_true")
    ap.add_argument("--cola", action="store_true")
    ap.add_argument("--aviso", action="store_true", help="mostrar el popup final también en modo manual")
    a = ap.parse_args()
    if a.cola:
        procesar_cola()
    elif a.codigo:
        r = preparar(a.codigo, a.estado, a.salida, not a.no_firmar)
        print(json.dumps(r, ensure_ascii=False, indent=2))
        if a.aviso:
            aviso_final(r)
    else:
        ap.print_help()
